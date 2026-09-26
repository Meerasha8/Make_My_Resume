import io
import json
import logging
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pydantic import BaseModel

from embedding_utils import cosine_similarity, embed_texts, relevance_scores
from llm import LLMUnavailable, chat_completion

logger = logging.getLogger(__name__)


class SkillGroup(BaseModel):
    category: str
    items: list[str] = []


class ExperienceItem(BaseModel):
    company: str
    role: str = ""
    duration: str = ""
    location: str = ""
    highlights: list[str] = []


class ProjectItem(BaseModel):
    name: str
    tech_stack: str = ""
    github_url: str = ""
    live_link: str = ""
    highlights: list[str] = []


class EducationItem(BaseModel):
    college_name: str = ""
    course_name: str = ""
    location: str = ""
    start_year: int | None = None
    end_year: int | None = None
    cgpa: float | None = None


class CertificateItem(BaseModel):
    certificate_name: str
    certificate_issuer: str = ""


class ResumeContent(BaseModel):
    education: list[EducationItem] = []
    skills: list[SkillGroup] = []
    experience: list[ExperienceItem] = []
    projects: list[ProjectItem] = []
    achievements: list[str] = []
    certificates: list[CertificateItem] = []


MAX_PROJECTS = 3
MAX_EXPERIENCE = 3
MAX_BULLETS = 3
MAX_ACHIEVEMENTS = 4
MAX_CERTIFICATES = 5

LLM_SCHEMA = {
    "skills": [{"category": "string", "items": ["string"]}],
    "experience": [{"company": "string (exact company_name from the data)", "highlights": ["string"]}],
    "projects": [{"name": "string (exact project name from the data)", "highlights": ["string"]}],
    "achievements": ["string"],
}

LLM_INSTRUCTIONS = (
    "You are writing a one-page ATS-friendly resume tailored to the job description, using ONLY the provided user data. "
    "Never invent employers, projects, skills, numbers or technologies that are not in the data.\n"
    f"- projects: pick the {MAX_PROJECTS} most relevant; use the exact project name from the data. Write exactly 3 bullets per "
    "project, in this order, without labels such as 'Problem:': (1) the problem the project solves and who it is for, "
    "(2) what was implemented - key features and technologies, (3) the result or impact. Use only results stated in the "
    "data; if none are given, describe the concrete outcome (e.g. deployed, used by, automated) without inventing numbers.\n"
    f"- experience: pick up to {MAX_EXPERIENCE} most relevant internships; use the exact company_name from the data.\n"
    f"- Write 2-{MAX_BULLETS} bullets for each experience. Every bullet is under 25 words and starts with a strong past-tense action verb.\n"
    "- In every bullet wrap the 1-3 most important technologies or outcomes in **double asterisks** so they render bold.\n"
    "- skills: group ONLY skill names that appear in the data into exactly these categories: Languages, "
    "Backend (frameworks, APIs), Frontend (web/mobile UI), Databases, Deployment & Tools (cloud, DevOps, Git, platforms), "
    "Core Concepts (DSA, OOP, DBMS, OS, networks). Put the most job-relevant skills first.\n"
    f"- achievements: up to {MAX_ACHIEVEMENTS} concise one-line achievements from the data, bolding key numbers or titles.\n"
    "Return strict raw JSON matching this schema, with no markdown and no prose: " + json.dumps(LLM_SCHEMA)
)


SKILL_CATEGORIES = ["Languages", "Backend", "Frontend", "Databases", "Deployment & Tools", "Core Concepts"]

# Common technologies (lower-case), so the grouping is consistent even when the AI is unavailable.
_SKILL_CATEGORY_MAP = {
    "Languages": "python, java, javascript, typescript, c, c++, c#, go, golang, rust, kotlin, swift, ruby, php, scala, r, dart, "
                 "sql, bash, shell scripting, solidity, matlab, perl, lua, haskell, elixir, zig, ocaml, julia, fortran, clojure, "
                 "erlang, f#, objective-c, assembly, groovy, cobol, visual basic",
    "Backend": "fastapi, django, flask, node, node.js, express, express.js, nestjs, spring, spring boot, .net, asp.net, laravel, "
               "ruby on rails, graphql, rest, rest api, rest apis, grpc, sqlalchemy, prisma, celery, kafka, rabbitmq, "
               "websockets, microservices, gin, echo, fiber, koa, hono, fastify, nest.js, phoenix, actix, axum, strapi, "
               "trpc, pydantic, socket.io, langchain, llamaindex",
    "Frontend": "react, react.js, reactjs, next.js, nextjs, vue, vue.js, angular, svelte, html, html5, css, css3, tailwind, "
                "tailwind css, bootstrap, sass, redux, jquery, react native, flutter, expo, material ui, shadcn, vite, solidjs, "
                "astro, remix, nuxt, gatsby, framer motion, jetpack compose, swiftui, three.js, d3.js, zustand, react query, "
                "tanstack query, ionic, webpack",
    "Databases": "postgresql, postgres, mysql, sqlite, mongodb, redis, mariadb, oracle, oracle db, dynamodb, cassandra, "
                 "firebase, firestore, supabase, neon, pgvector, elasticsearch, sql server, microsoft sql server, mssql, neo4j, "
                 "couchbase, influxdb, cockroachdb, planetscale, clickhouse, snowflake, bigquery, pinecone, chromadb, qdrant, "
                 "weaviate, milvus, faiss",
    "Deployment & Tools": "docker, kubernetes, k8s, aws, gcp, azure, google cloud, git, github, gitlab, bitbucket, linux, ci/cd, "
                          "github actions, jenkins, vercel, render, netlify, heroku, railway, nginx, terraform, ansible, "
                          "postman, jira, figma, vs code, cloudflare, digitalocean, podman, argocd, grafana, prometheus, "
                          "fly.io, circleci, gitlab ci, helm, openshift, aws lambda, ec2, s3, cloud run, npm, yarn, pnpm, maven, "
                          "gradle, selenium, jest, pytest, docker compose",
    "Core Concepts": "data structures, algorithms, data structures & algorithms, data structures and algorithms, dsa, oop, "
                     "object oriented programming, dbms, operating systems, os, computer networks, cn, system design, "
                     "design patterns, machine learning, deep learning, computer architecture, software engineering, agile, "
                     "distributed systems, concurrency, multithreading, graph theory, cryptography, compilers, compiler design, "
                     "cloud computing",
}
_SKILL_LOOKUP = {name.strip(): category for category, names in _SKILL_CATEGORY_MAP.items() for name in names.split(",")}


def _canonical_category(name: str | None) -> str | None:
    return next((category for category in SKILL_CATEGORIES if category.casefold() == (name or "").strip().casefold()), None)


def categorize_skill(skill: str, suggested: str | None = None) -> str:
    """Dictionary match, then name heuristics (…DB/SQL -> Databases, …UI -> Frontend), then the AI's suggestion if it is
    one of the fixed categories, then the nearest dictionary skills by meaning."""
    key = skill.strip().casefold()
    for candidate in (key, key.removesuffix(".js"), key.replace(".js", "js")):
        if candidate in _SKILL_LOOKUP:
            return _SKILL_LOOKUP[candidate]
    if "algorithm" in key or "data structure" in key:
        return "Core Concepts"
    if re.search(r"db\b|sql|database", key):
        return "Databases"
    if re.search(r"\bui\b", key):
        return "Frontend"
    if _canonical_category(suggested):
        return _canonical_category(suggested)
    return _nearest_known_category(skill)


def _nearest_known_category(skill: str, k: int = 3) -> str:
    """Vote among the k dictionary skills closest in meaning (measured 17/30 on unseen skills vs 9/30 for
    comparing against category descriptions)."""
    known = list(_SKILL_LOOKUP)
    reference = embed_texts([f"Technical skill: {name}" for name in known])
    query = embed_texts([f"Technical skill: {skill}"])[0]
    nearest = sorted(((cosine_similarity(query, vector), _SKILL_LOOKUP[name]) for name, vector in zip(known, reference)), reverse=True)[:k]
    votes: dict[str, float] = {}
    for similarity, category in nearest:
        votes[category] = votes.get(category, 0.0) + similarity
    return max(votes, key=votes.get)


def group_skills(skills: list[tuple[str, str | None]]) -> list[SkillGroup]:
    """Group (skill, suggested category) pairs into the fixed categories, keeping the given order within each."""
    grouped: dict[str, list[str]] = {category: [] for category in SKILL_CATEGORIES}
    for skill, suggested in skills:
        items = grouped[categorize_skill(skill, suggested)]
        if skill not in items:
            items.append(skill)
    return [SkillGroup(category=category, items=items) for category, items in grouped.items() if items]


def _clean_bullets(values: Any, limit: int = MAX_BULLETS) -> list[str]:
    bullets = []
    for value in values or []:
        # Strip list markers ("•", "-", "* ") but keep a leading **bold** span intact.
        text = re.sub(r"^(?:[•\-–]|\*(?!\*))\s*", "", str(value).strip()).strip()
        if text:
            bullets.append(text)
    return bullets[:limit]


def _sentences(text: str | None) -> list[str]:
    return [part.strip() for part in re.split(r"(?<=[.!?])\s+", text or "") if part.strip()]


class ResumeService:
    # ------------------------------------------------------------------ content selection

    def _filter_user_data_with_rag(self, job_description: str, user_data: dict[str, Any]) -> dict[str, Any]:
        def rank(items, text_extractor, top_k):
            items = list(items or [])
            scores = relevance_scores(job_description, [text_extractor(item) for item in items])
            ranked = sorted(zip(scores, range(len(items))), key=lambda pair: pair[0], reverse=True)
            return [items[index] for _, index in ranked[:top_k]]

        return {
            "user_details": user_data.get("user_details", {}),
            "education": user_data.get("education", []),
            # "Technical skill:" gives bare skill names enough context to embed well (measured: 14/17 vs 12/17 bare).
            "skills": rank(user_data.get("skills"), lambda x: f"Technical skill: {x.get('name', '')}. {x.get('description') or ''}", 30),
            "internship": rank(user_data.get("internship"), lambda x: f"{x.get('role', '')} {x.get('description', '')}", 5),
            "projects": rank(
                user_data.get("projects"),
                lambda x: f"{x.get('name', '')} {x.get('description', '')} {x.get('tech_stack', '')}",
                6,
            ),
            "achievements": rank(user_data.get("achievements"), lambda x: x.get("description", ""), 8),
            "certificates": rank(
                user_data.get("certificates"),
                lambda x: f"{x.get('certificate_name', '')} {x.get('certificate_issuer', '')}",
                MAX_CERTIFICATES,
            ),
        }

    def _ask_llm(self, job_description: str, filtered_data: dict[str, Any]) -> dict[str, Any] | None:
        llm_data = {key: value for key, value in filtered_data.items() if key not in ("user_details", "education", "certificates")}
        messages = [
            {"role": "system", "content": "Return only raw valid JSON matching the requested schema. Do not use markdown codeblocks."},
            {
                "role": "user",
                "content": (
                    f"Job description:\n{job_description}\n\n"
                    f"User data:\n{json.dumps(llm_data, ensure_ascii=False, indent=2)}\n\n"
                    f"Instructions:\n{LLM_INSTRUCTIONS}"
                ),
            },
        ]
        for attempt in range(3):
            try:
                content = chat_completion(messages, max_tokens=4000, json_mode=True).strip()
                if content.startswith("```"):
                    content = content.strip("`").removeprefix("json").strip()
                return json.loads(content)
            except LLMUnavailable:
                return None
            except Exception as e:
                logger.warning("Groq resume content attempt %s failed: %s", attempt, e)
        return None

    def select_resume_content(self, job_description: str, user_data: dict[str, Any]) -> ResumeContent:
        filtered = self._filter_user_data_with_rag(job_description, user_data)
        draft = self._ask_llm(job_description, filtered)
        return self._build_content(draft or {}, filtered)

    def _build_content(self, draft: dict[str, Any], data: dict[str, Any]) -> ResumeContent:
        """Merge the AI draft with the user's real data. Facts (names, dates, links, stacks) always come from the data;
        anything in the draft that doesn't match the data is dropped, and empty sections fall back to the raw data."""
        # Experience: match on company name, keep role/duration from the data.
        remaining_internships = list(data.get("internship", []))
        experience = []
        for item in draft.get("experience") or []:
            if not isinstance(item, dict):
                continue
            company = str(item.get("company", "")).strip().casefold()
            source = next((row for row in remaining_internships if (row.get("company_name") or "").strip().casefold() == company), None)
            if source:
                remaining_internships.remove(source)
                experience.append(self._experience_from(source, _clean_bullets(item.get("highlights"))))
        if not experience:
            experience = [self._experience_from(row) for row in data.get("internship", [])]

        # Projects: match on project name, keep stack and links from the data.
        projects_by_name = {(row.get("name") or "").strip().casefold(): row for row in data.get("projects", [])}
        projects = []
        for item in draft.get("projects") or []:
            if not isinstance(item, dict):
                continue
            source = projects_by_name.pop(str(item.get("name", "")).strip().casefold(), None)
            if source:
                projects.append(self._project_from(source, _clean_bullets(item.get("highlights"))))
        if not projects:
            projects = [self._project_from(row) for row in data.get("projects", [])]

        # Skills: keep only skills the user actually has, in the AI's relevance order, grouped into the fixed categories.
        known_skills = {(row.get("name") or "").strip().casefold(): row["name"].strip() for row in data.get("skills", []) if row.get("name")}
        selected: list[tuple[str, str | None]] = []
        for group in draft.get("skills") or []:
            if not isinstance(group, dict):
                continue
            for skill in group.get("items") or []:
                name = known_skills.get(str(skill).strip().casefold())
                if name:
                    selected.append((name, str(group.get("category") or "")))
        if not selected:
            selected = [(name, None) for name in known_skills.values()]
        skills = group_skills(selected)

        achievements = _clean_bullets(draft.get("achievements"), MAX_ACHIEVEMENTS) or _clean_bullets(
            [row.get("description") for row in data.get("achievements", [])], MAX_ACHIEVEMENTS
        )

        education = sorted(
            (EducationItem(**{key: row.get(key) for key in EducationItem.model_fields if row.get(key) is not None}) for row in data.get("education", [])),
            key=lambda item: item.end_year or 0,
            reverse=True,
        )
        certificates = [
            CertificateItem(certificate_name=row["certificate_name"], certificate_issuer=row.get("certificate_issuer") or "")
            for row in data.get("certificates", [])
            if row.get("certificate_name")
        ]

        return ResumeContent(
            education=education,
            skills=skills,
            experience=experience[:MAX_EXPERIENCE],
            projects=projects[:MAX_PROJECTS],
            achievements=achievements,
            certificates=certificates[:MAX_CERTIFICATES],
        )

    @staticmethod
    def _experience_from(row: dict[str, Any], highlights: list[str] | None = None) -> ExperienceItem:
        return ExperienceItem(
            company=row.get("company_name") or "",
            role=row.get("role") or "",
            duration=row.get("duration") or "",
            highlights=highlights or _sentences(row.get("description"))[:MAX_BULLETS],
        )

    @staticmethod
    def _project_from(row: dict[str, Any], highlights: list[str] | None = None) -> ProjectItem:
        return ProjectItem(
            name=row.get("name") or "",
            tech_stack=row.get("tech_stack") or "",
            github_url=row.get("github_url") or "",
            live_link=row.get("live_link") or "",
            highlights=highlights or _sentences(row.get("description"))[:MAX_BULLETS],
        )

    # ------------------------------------------------------------------ DOCX rendering (single-page ATS layout)

    FONT = "Cambria"
    BODY_SIZE = Pt(10)
    TEXT_WIDTH = Inches(7.5)  # US Letter minus 0.5" side margins
    SKILL_LABEL_WIDTH = Inches(1.6)

    def render_docx(self, content: ResumeContent, user_details: dict[str, Any] | None = None) -> bytes:
        details = user_details or {}
        document = Document()

        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.45)
        section.left_margin = section.right_margin = Inches(0.5)

        normal = document.styles["Normal"]
        normal.font.name = self.FONT
        normal.font.size = self.BODY_SIZE
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), self.FONT)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.0
        document.core_properties.title = f"{details.get('name') or 'Resume'} - Resume"

        self._render_header(document, details)

        if content.education:
            self._section_heading(document, "Education")
            for item in content.education:
                years = " – ".join(str(year) for year in (item.start_year, item.end_year) if year)
                self._two_column(document, item.college_name, years, bold=True)
                self._two_column(document, item.course_name, f"CGPA: {item.cgpa:g}" if item.cgpa else item.location, italic=True, size=Pt(9.5))

        if content.skills:
            self._section_heading(document, "Technical Skills")
            for group in content.skills:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.tab_stops.add_tab_stop(self.SKILL_LABEL_WIDTH)
                paragraph.paragraph_format.left_indent = self.SKILL_LABEL_WIDTH
                paragraph.paragraph_format.first_line_indent = -self.SKILL_LABEL_WIDTH
                paragraph.add_run(group.category).bold = True
                paragraph.add_run("\t" + ", ".join(group.items))

        if content.experience:
            self._section_heading(document, "Experience")
            for index, item in enumerate(content.experience):
                self._two_column(document, item.company, item.duration, bold=True, space_before=Pt(3) if index else None)
                if item.role or item.location:
                    self._two_column(document, item.role, item.location, italic=True, size=Pt(9.5))
                self._bullets(document, item.highlights)

        if content.projects:
            self._section_heading(document, "Projects")
            for index, item in enumerate(content.projects):
                paragraph = self._row(document, space_before=Pt(3) if index else None)
                name = paragraph.add_run(item.name)
                name.bold = True
                name.font.size = Pt(10.5)
                if item.tech_stack:
                    paragraph.add_run("  |  ")
                    stack = paragraph.add_run(item.tech_stack)
                    stack.italic = True
                    stack.font.size = Pt(9.5)
                links = [(label, url) for label, url in (("Live", item.live_link), ("GitHub", item.github_url)) if url]
                if links:
                    paragraph.add_run("\t")
                    for link_index, (label, url) in enumerate(links):
                        if link_index:
                            paragraph.add_run("  |  ")
                        self._hyperlink(paragraph, url, label)
                self._bullets(document, item.highlights)

        if content.achievements:
            self._section_heading(document, "Achievements")
            self._bullets(document, content.achievements)

        if content.certificates:
            self._section_heading(document, "Certifications")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for index, item in enumerate(content.certificates):
                if index:
                    paragraph.add_run("  |  ")
                paragraph.add_run(item.certificate_name).bold = True
                if item.certificate_issuer:
                    paragraph.add_run(f" – {item.certificate_issuer}")

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _render_header(self, document, details: dict[str, Any]) -> None:
        name = document.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name.add_run((details.get("name") or "Your Name").upper())
        run.bold = True
        run.font.size = Pt(18)

        contact = [(details.get("phone"), None), (details.get("email"), f"mailto:{details.get('email')}"), (details.get("location"), None)]
        self._separated_line(document, [(text, url) for text, url in contact if text])

        links = [
            (self._link_label("Linkedin", details.get("linkedin")), details.get("linkedin")),
            (self._link_label("Github", details.get("github")), details.get("github")),
            ("Portfolio", details.get("portfolio")),
        ]
        self._separated_line(document, [(text, url) for text, url in links if url])

    def _separated_line(self, document, parts: list[tuple[str, str | None]]) -> None:
        if not parts:
            return
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for index, (text, url) in enumerate(parts):
            if index:
                paragraph.add_run("  |  ").font.size = Pt(9.5)
            if url:
                self._hyperlink(paragraph, url, text, size=Pt(9.5))
            else:
                paragraph.add_run(text).font.size = Pt(9.5)

    @staticmethod
    def _link_label(label: str, url: str | None) -> str:
        path = re.sub(r"^[a-z]+://", "", (url or "").strip().rstrip("/"), flags=re.I)
        handle = path.rsplit("/", 1)[-1] if "/" in path else ""
        return f"{label}/{handle}" if handle else label

    def _section_heading(self, document, title: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(7)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(title)
        run.bold = True
        run.font.size = Pt(11.5)

        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
            bottom.set(qn(key), value)
        border.append(bottom)
        paragraph._p.get_or_add_pPr().insert_element_before(
            border, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr", "w:pPrChange"
        )

    def _row(self, document, space_before=None):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.tab_stops.add_tab_stop(self.TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
        if space_before is not None:
            paragraph.paragraph_format.space_before = space_before
        return paragraph

    def _two_column(self, document, left: str, right: str, *, bold=False, italic=False, size=None, space_before=None) -> None:
        paragraph = self._row(document, space_before)
        for text in (left or "", "\t" + right if right else ""):
            if text:
                run = paragraph.add_run(text)
                run.bold, run.italic = bold, italic
                if size:
                    run.font.size = size

    def _bullets(self, document, items: list[str]) -> None:
        for text in items:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.first_line_indent = Inches(-0.15)
            # **text** from the AI draft renders bold, matching the reference layout.
            for part in re.split(r"(\*\*[^*]+\*\*)", text):
                if part.startswith("**") and part.endswith("**") and len(part) > 4:
                    paragraph.add_run(part[2:-2]).bold = True
                elif part:
                    paragraph.add_run(part.replace("**", ""))

    @staticmethod
    def _hyperlink(paragraph, url: str, text: str, size=None) -> None:
        if not re.match(r"^[a-z]+:", url, flags=re.I):
            url = "https://" + url
        relationship_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        paragraph._p.append(hyperlink)
        run = paragraph.add_run(text)
        if size:
            run.font.size = size
        hyperlink.append(run._r)  # moves the run inside the hyperlink element
